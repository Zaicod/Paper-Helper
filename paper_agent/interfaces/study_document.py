"""Deterministic Word serialization from a saved structured report only."""
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from paper_agent.core.study import StudyReport
from paper_agent.interfaces.study_citations import StudyCitationFormatter

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def render_study_document(report: StudyReport) -> bytes:
    report = StudyReport.model_validate(report)
    citations = StudyCitationFormatter(report)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
        style.font.size = Pt(11 if name in ('Normal', 'List Bullet') else
                             {'Title': 24, 'Subtitle': 11, 'Heading 1': 18, 'Heading 2': 14, 'Heading 3': 12}[name])
        style.paragraph_format.space_after = Pt(7)
    document.styles['Normal'].paragraph_format.line_spacing = 1.35
    document.add_heading(report.topic + ' · 学习报告', 0)
    document.add_paragraph('生成时间：' + report.generated_at.isoformat(timespec='seconds'))
    document.add_paragraph('依据：开放 PDF 的带页码证据片段。学习体会与计划为 Agent 建议，请结合原文核对。')
    if report.mode == 'demo':
        document.add_paragraph('离线演示报告：仅用于流程演示，不代表真实论文分析。')
    document.add_heading('论文清单', 1)
    for index, paper in enumerate(report.papers, 1):
        document.add_paragraph(f'{index}. {paper.title}\n{paper.page_url}')
    for index, (paper, item) in enumerate(zip(report.papers, report.analyses), 1):
        def readable(text):
            return citations.format(text, paper.paper_id)
        document.add_page_break()
        document.add_heading(f'{index}. {item.title}', 1)
        document.add_heading('论文概述', 2)
        document.add_paragraph(readable(item.overview))
        document.add_heading('1. 研究背景与问题', 2)
        document.add_paragraph(readable(item.background))
        document.add_heading('2. 模型方法（模块解释）', 2)
        for module in item.methods:
            document.add_heading(readable(module.name), 3)
            document.add_paragraph('作用：' + readable(module.purpose))
            document.add_paragraph('输入：' + readable(module.inputs))
            document.add_paragraph('输出：' + readable(module.outputs))
        document.add_heading('3. 创新点与实验结果', 2)
        document.add_heading('创新点', 3)
        for value in item.innovations:
            document.add_paragraph(readable(value), style='List Bullet')
        document.add_heading('实验结果', 3)
        for value in item.results:
            document.add_paragraph(readable(value), style='List Bullet')
        document.add_heading('4. 学习体会（Agent 建议）', 2)
        for value in item.reflection:
            document.add_paragraph(readable(value), style='List Bullet')
        document.add_heading('证据定位', 2)
        seen = set()
        for reference in item.citations:
            if reference.chunk_id not in seen:
                number = citations.by_paper[paper.paper_id][reference.chunk_id]
                document.add_paragraph(f'[{number}] 第 {reference.page_number} 页\n{reference.source_url}')
                seen.add(reference.chunk_id)
    document.add_page_break()
    document.add_heading('跨论文总结' if len(report.papers) > 1 else '方法总结', 1)
    for value in report.synthesis:
        document.add_paragraph(citations.format(value), style='List Bullet')
    document.add_heading('建议学习计划', 1)
    for value in report.learning_plan:
        document.add_paragraph(citations.format(value), style='List Bullet')
    if citations.unresolved:
        document.add_paragraph('标注“引用待核对”的位置没有对应的已确认原文证据，请核对原论文。')
    document.add_heading('参考论文', 1)
    for paper in report.papers:
        document.add_paragraph(f'{paper.title}\n作者：{"、".join(paper.authors)}\n{paper.page_url}')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    target = BytesIO()
    document.save(target)
    return target.getvalue()
